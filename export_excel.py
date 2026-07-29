"""
Export database contents to an Excel workbook matching the layout of stats-4-23.xlsx.

Called from the Flask app (/admin/export) or standalone:
    python export_excel.py [output_path.xlsx]
"""

import os
import sys
import io
from datetime import datetime

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app import app, db
from models import Category, Metric, Branch, Entry, EntryValue

MONTHS = ['January','February','March','April','May','June',
          'July','August','September','October','November','December']

# ── Exact column order matching stats-4-23.xlsx ───────────────────────────────

BRANCH_STATS_COLUMNS = [
    'Month Num', 'Month', 'BRANCH',
    'New Library Card Registrations, Adult (includes YA)',
    'New Library Card Registrations, Juvenile',
    'Gate Count',
    'PC Reservations',
    'WiFi - Unique Sessions',
    'External Party Library Room Use',
    'Total Branch Circulation',
    'Hotspots Circulation',
    'Curbside',
    'ILL - Sent (Main ONLY)',
    'ILL - Received (Main ONLY)',
    'ICLs - Sent (MAIN ONLY)',
    'ICLs - Received (MAIN ONLY)',
    'Total Prints per Month',
    'I2:  ONSITE Sessions 0-5',
    'I3:   ONSITE Sessions 6-11',
    'I4: ONSITE Sessions 12-18',
    'I5:   ONSITE Sessions 19+',
    'I6:  ONSITE Sessions GENERAL INTEREST',
    'ONSITE Attendance 0-5',
    'ONSITE Attendance 6-11',
    'ONSITE Attendance 12-18',
    'ONSITE Attendance 19+',
    'ONSITE Attendance General Interest',
    'OFFSITE Sessions 0-5',
    'OFFSITE Sessions 6-11',
    'OFFSITE Sessions 12-18',
    'OFFSITE Sessions 19+',
    'OFFSITE Sessions General Interest',
    'OFFSITE Attendance 0-5',
    'OFFSITE Attendance 6-11',
    'OFFSITE Attendance 12-18',
    'OFFSITE Attendance 19+',
    'OFFSITE Attendance General Interest',
    'VIRTUAL Sessions 0-5',
    'VIRTUAL Sessions 6-11',
    'VIRTUAL Sessions 12-18',
    'VIRTUAL Sessions 19+',
    'VIRTUAL Sessions General Interest',
    'VIRTUAL Attendance 0-5',
    'VIRTUAL Attendance 6-11',
    'VIRTUAL Attendance 12-18',
    'VIRTUAL Attendance 19+',
    'VIRTUAL Attendance General Interest',
    'I21: NUMBER OF OUTREACH ACTIVITIES Conducted',
    'Outreach Attendance (YCL Internal)',
    'I22: TOTAL # TAKE & MAKES and OTHER PASSIVE PROGRAM PARTICIPANTS\n',
    'I23: NUMBER OF STAFF TAKING TRAINING',
    'I24: NUMBER OF HOURS STAFF ATTENDED TRAINING',
    '1-on-1 Total for Month',
    'Email Address',
    'Year',
    'Locker Circulation',
]

# Maps Excel column name → DB metric name (for Branch Stats)
BRANCH_STATS_MAP = {
    'New Library Card Registrations, Adult (includes YA)': 'New Library Card Registrations, Adult',
    'New Library Card Registrations, Juvenile':            'New Library Card Registrations, Juvenile',
    'Gate Count':                                          'Gate Count',
    'PC Reservations':                                     'PC Reservations',
    'WiFi - Unique Sessions':                              'WiFi - Unique Sessions',
    'External Party Library Room Use':                     'External Party Library Room Use',
    'Total Branch Circulation':                            'Total Branch Circulation',
    'Hotspots Circulation':                                'Hotspots Circulation',
    'Curbside':                                            'Curbside',
    'ILL - Sent (Main ONLY)':                              'ILL - Sent (Main ONLY)',
    'ILL - Received (Main ONLY)':                          'ILL - Received (Main ONLY)',
    'ICLs - Sent (MAIN ONLY)':                             'ICLs - Sent (Main ONLY)',
    'ICLs - Received (MAIN ONLY)':                         'ICLs - Received (Main ONLY)',
    'Total Prints per Month':                              'Total Prints per Month',
    'I2:  ONSITE Sessions 0-5':                            'ONSITE Sessions 0-5',
    'I3:   ONSITE Sessions 6-11':                          'ONSITE Sessions 6-11',
    'I4: ONSITE Sessions 12-18':                           'ONSITE Sessions 12-18',
    'I5:   ONSITE Sessions 19+':                           'ONSITE Sessions 19+',
    'I6:  ONSITE Sessions GENERAL INTEREST':               'ONSITE Sessions General Interest',
    'ONSITE Attendance 0-5':                               'ONSITE Attendance 0-5',
    'ONSITE Attendance 6-11':                              'ONSITE Attendance 6-11',
    'ONSITE Attendance 12-18':                             'ONSITE Attendance 12-18',
    'ONSITE Attendance 19+':                               'ONSITE Attendance 19+',
    'ONSITE Attendance General Interest':                  'ONSITE Attendance General Interest',
    'OFFSITE Sessions 0-5':                                'OFFSITE Sessions 0-5',
    'OFFSITE Sessions 6-11':                               'OFFSITE Sessions 6-11',
    'OFFSITE Sessions 12-18':                              'OFFSITE Sessions 12-18',
    'OFFSITE Sessions 19+':                                'OFFSITE Sessions 19+',
    'OFFSITE Sessions General Interest':                   'OFFSITE Sessions General Interest',
    'OFFSITE Attendance 0-5':                              'OFFSITE Attendance 0-5',
    'OFFSITE Attendance 6-11':                             'OFFSITE Attendance 6-11',
    'OFFSITE Attendance 12-18':                            'OFFSITE Attendance 12-18',
    'OFFSITE Attendance 19+':                              'OFFSITE Attendance 19+',
    'OFFSITE Attendance General Interest':                 'OFFSITE Attendance General Interest',
    'VIRTUAL Sessions 0-5':                                'VIRTUAL Sessions 0-5',
    'VIRTUAL Sessions 6-11':                               'VIRTUAL Sessions 6-11',
    'VIRTUAL Sessions 12-18':                              'VIRTUAL Sessions 12-18',
    'VIRTUAL Sessions 19+':                                'VIRTUAL Sessions 19+',
    'VIRTUAL Sessions General Interest':                   'VIRTUAL Sessions General Interest',
    'VIRTUAL Attendance 0-5':                              'VIRTUAL Attendance 0-5',
    'VIRTUAL Attendance 6-11':                             'VIRTUAL Attendance 6-11',
    'VIRTUAL Attendance 12-18':                            'VIRTUAL Attendance 12-18',
    'VIRTUAL Attendance 19+':                              'VIRTUAL Attendance 19+',
    'VIRTUAL Attendance General Interest':                 'VIRTUAL Attendance General Interest',
    'I21: NUMBER OF OUTREACH ACTIVITIES Conducted':        'Number of Outreach Activities Conducted',
    'Outreach Attendance (YCL Internal)':                  'Outreach Attendance',
    'I22: TOTAL # TAKE & MAKES and OTHER PASSIVE PROGRAM PARTICIPANTS\n':
                                                           'Take & Makes / Other Passive Program Participants',
    'I23: NUMBER OF STAFF TAKING TRAINING':                'Number of Staff Taking Training',
    'I24: NUMBER OF HOURS STAFF ATTENDED TRAINING':        'Number of Hours Staff Attended Training',
    '1-on-1 Total for Month':                              '1-on-1 Total for Month',
    'Locker Circulation':                                  'Locker Circulation',
}

ONLINE_STATS_COLUMNS = [
    'Month Num', 'Month',
    'yclibrary.org - web sessions',
    'ychistory.org - views',
    'patchworktales.org  - views',
    'Dial A Story - CALLS',
    'Dial A Story - VIEWS',
    'DSpace - Views',
    'Beanstack - Sessions',
    'LibraryCalendar - Sessions',
    'LibGuides - Sessions',
    'DigitalLearn.org - Sessions',
    'DigitalLearn.org - Completed Courses',
    'LOTE4Kids - Stories Watched',
    'LOTE4Kids - Actvitities',
    'LOTE4Kids - Logins',
    'Youtube - Subscribers',
    'YouTube - Views',
    'YouTube - Hours Watched',
    'YCL News - Subscriber',
    'Website Messages',
    'YCL - App - Users',
    'YCL - App - Sessions',
    'Facebook Followers',
    'Instragram - Subscribers',
    'YouTube Uploads',
    'Dial A Story Uploads',
    'Year',
]

ONLINE_STATS_MAP = {
    'yclibrary.org - web sessions':         'yclibrary.org - Web Sessions',
    'ychistory.org - views':                'ychistory.org - Views',
    'patchworktales.org  - views':          'patchworktales.org - Views',
    'Dial A Story - CALLS':                 'Dial A Story - Calls',
    'Dial A Story - VIEWS':                 'Dial A Story - Views',
    'DSpace - Views':                       'DSpace - Views',
    'Beanstack - Sessions':                 'Beanstack - Sessions',
    'LibraryCalendar - Sessions':           'LibraryCalendar - Sessions',
    'LibGuides - Sessions':                 'LibGuides - Sessions',
    'DigitalLearn.org - Sessions':          'DigitalLearn.org - Sessions',
    'DigitalLearn.org - Completed Courses': 'DigitalLearn.org - Completed Courses',
    'LOTE4Kids - Stories Watched':          'LOTE4Kids - Stories Watched',
    'LOTE4Kids - Actvitities':              'LOTE4Kids - Activities',
    'LOTE4Kids - Logins':                   'LOTE4Kids - Logins',
    'Youtube - Subscribers':                'YouTube - Subscribers',
    'YouTube - Views':                      'YouTube - Views',
    'YouTube - Hours Watched':              'YouTube - Hours Watched',
    'YCL News - Subscriber':                'YCL News - Subscribers',
    'Website Messages':                     'Website Messages',
    'YCL - App - Users':                    'YCL App - Users',
    'YCL - App - Sessions':                 'YCL App - Sessions',
    'Facebook Followers':                   'Facebook Followers',
    'Instragram - Subscribers':             'Instagram - Subscribers',
    'YouTube Uploads':                      'YouTube Uploads',
    'Dial A Story Uploads':                 'Dial A Story Uploads',
}

QRTLY_COLUMNS = [
    'Timestamp', 'Quarter', 'Branch or Location',
    'Total # of Transactions for the Week', 'Year',
]

_QUARTER_LABELS = {
    1: 'Quarter 1',
    2: 'Quarter 2',
    3: 'Quarter 3',
    4: 'Quarter 4',
}

# ── Styling helpers ───────────────────────────────────────────────────────────

def style_header_row(ws, n_cols):
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E79')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = 'A2'

# ── Sheet writers ─────────────────────────────────────────────────────────────

def write_branch_stats(wb, cat):
    ws = wb.create_sheet('Branch Stats')
    ws.append(BRANCH_STATS_COLUMNS)

    # Build metric_name → metric_id lookup
    metric_by_name = {m.name: m.id for m in cat.metrics}
    # Excel column → metric_id
    col_to_metric_id = {
        col: metric_by_name[db_name]
        for col, db_name in BRANCH_STATS_MAP.items()
        if db_name in metric_by_name
    }

    entries = (Entry.query
               .filter_by(category_id=cat.id)
               .order_by(Entry.year, Entry.month, Entry.branch_id)
               .all())

    for entry in entries:
        # Build metric_id → value map for this entry
        val_map = {ev.metric_id: ev.value_number
                   for ev in entry.values if ev.value_number is not None}

        branch_name = entry.branch.name.upper() if entry.branch else ''
        month_name  = MONTHS[entry.month - 1] if entry.month else ''

        row = []
        for col in BRANCH_STATS_COLUMNS:
            if col == 'Month Num':
                row.append(entry.month)
            elif col == 'Month':
                row.append(month_name)
            elif col == 'BRANCH':
                row.append(branch_name)
            elif col == 'Year':
                row.append(entry.year)
            elif col == 'Email Address':
                row.append(None)
            elif col in col_to_metric_id:
                row.append(val_map.get(col_to_metric_id[col]))
            else:
                row.append(None)
        ws.append(row)

    style_header_row(ws, len(BRANCH_STATS_COLUMNS))
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 22
    return len(entries)


def write_online_stats(wb, cat):
    ws = wb.create_sheet('Online Stats')
    ws.append(ONLINE_STATS_COLUMNS)

    metric_by_name = {m.name: m.id for m in cat.metrics}
    col_to_metric_id = {
        col: metric_by_name[db_name]
        for col, db_name in ONLINE_STATS_MAP.items()
        if db_name in metric_by_name
    }

    entries = (Entry.query
               .filter_by(category_id=cat.id)
               .order_by(Entry.year, Entry.month)
               .all())

    for entry in entries:
        val_map = {ev.metric_id: ev.value_number
                   for ev in entry.values if ev.value_number is not None}
        month_name = MONTHS[entry.month - 1] if entry.month else ''

        row = []
        for col in ONLINE_STATS_COLUMNS:
            if col == 'Month Num':
                row.append(entry.month)
            elif col == 'Month':
                row.append(month_name)
            elif col == 'Year':
                row.append(entry.year)
            elif col in col_to_metric_id:
                row.append(val_map.get(col_to_metric_id[col]))
            else:
                row.append(None)
        ws.append(row)

    style_header_row(ws, len(ONLINE_STATS_COLUMNS))
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 12
    return len(entries)


def write_quarterly_ref(wb, cat):
    ws = wb.create_sheet('Qrtly Ref Stats')
    ws.append(QRTLY_COLUMNS)

    metric = next((m for m in cat.metrics
                   if m.name == 'Total Transactions for the Week'), None)

    entries = (Entry.query
               .filter_by(category_id=cat.id)
               .order_by(Entry.year, Entry.quarter, Entry.branch_id)
               .all())

    count = 0
    for entry in entries:
        if not entry.quarter:
            continue
        val_map = {ev.metric_id: ev.value_number for ev in entry.values}
        val = val_map.get(metric.id) if metric else None

        branch_name = entry.branch.name if entry.branch else ''
        quarter_label = _QUARTER_LABELS.get(entry.quarter, f'Quarter {entry.quarter}')

        ws.append([
            datetime.now(),     # Timestamp placeholder
            quarter_label,
            branch_name,
            val,
            entry.year,
        ])
        count += 1

    style_header_row(ws, len(QRTLY_COLUMNS))
    ws.column_dimensions['C'].width = 28
    return count


# ── Main entry point ──────────────────────────────────────────────────────────

def generate_export(output=None):
    """
    Generate the export workbook.
    output: file path string or file-like object (BytesIO).
    Returns a BytesIO if no output given.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default empty sheet

    from flask import has_app_context
    ctx = None if has_app_context() else app.app_context()
    if ctx:
        ctx.push()
    try:
        totals = {}

        cat = Category.query.filter_by(name='Branch Stats').first()
        if cat:
            totals['Branch Stats'] = write_branch_stats(wb, cat)

        cat = Category.query.filter_by(name='Online Stats').first()
        if cat:
            totals['Online Stats'] = write_online_stats(wb, cat)

        cat = Category.query.filter_by(name='Quarterly Reference Stats').first()
        if cat:
            totals['Quarterly Ref Stats'] = write_quarterly_ref(wb, cat)
    finally:
        if ctx:
            ctx.pop()

    for sheet, count in totals.items():
        print(f"  {sheet}: {count} rows exported")

    if output is None:
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    wb.save(output)
    return output


if __name__ == '__main__':
    path = sys.argv[1] if len(sys.argv) > 1 else f'library_stats_export_{datetime.now():%Y%m%d}.xlsx'
    print(f"Exporting to {path} ...")
    generate_export(path)
    print("Done!")
