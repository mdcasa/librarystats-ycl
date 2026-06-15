"""
Extract legacy programming data from OldConversion .xls files into a backup .xlsx spreadsheet.
Covers FY2015-16 through FY2018-19 (pre-DB era for programming data).
Run from repo root: python3 extract_legacy_programs.py
"""

import xlrd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = "Data files/OldConversion/"
OUT  = "Data files/historical_data/Legacy_Programming_FY2015-2019.xlsx"

MONTHS = ['July', 'Aug', 'Sept', 'Oct', 'Nov', 'Dec', 'Jan', 'Feb', 'Mar', 'April', 'May', 'June']
# col 3 = July, col 4 = Aug, ... col 14 = June
MONTH_COL_START = 3

BRANCH_RENAMES = {
    'Rock Hill (SDD Site)': 'Rock Hill',
    'Bookmobile/Outreach':  'Bookmobile/Outreach',
    'Rock Hill':            'Rock Hill',
    'Clover':               'Clover',
    'Fort Mill':            'Fort Mill',
    'Lake Wylie':           'Lake Wylie',
    'York':                 'York',
}

# ─────────────────────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────────────────────

def cell(ws, r, c):
    v = ws.cell_value(r, c)
    if isinstance(v, float) and v == int(v):
        return int(v)
    return v if str(v).strip() else None


def month_from_col(c):
    idx = c - MONTH_COL_START
    return MONTHS[idx] if 0 <= idx < 12 else None


def fy_label(fy_start):
    return f"FY{fy_start}-{str(fy_start + 1)[2:]}"


def year_for_month(fy_start, month):
    """Return calendar year for a given fiscal month (July=fy_start, Jan-June=fy_start+1)."""
    if month in ('July', 'Aug', 'Sept', 'Oct', 'Nov', 'Dec'):
        return fy_start
    return fy_start + 1


# ─────────────────────────────────────────────────────────────────────────────
# Parser: per-branch Programs / Passive Programs sheets (FY16-17 through FY18-19)
# ─────────────────────────────────────────────────────────────────────────────

def parse_section_label(label):
    """Return (age_group, stat_type) or None if not a recognized data section."""
    label = label.strip()
    age_map = {
        '0-5':        ['Age 0-5'],
        '6-11':       ['Age 6-11'],
        '12-18':      ['Age 12-18'],
        'Mixed Ages': ['Mixed Ages'],
        '18+':        ['Age 18+'],
    }
    for age, prefixes in age_map.items():
        for prefix in prefixes:
            if label.startswith(prefix):
                rest = label[len(prefix):].strip()
                if 'Session' in rest or rest == '':
                    return (age, 'Sessions')
                elif 'Particip' in rest:
                    return (age, 'Participation')
    return None


def parse_branch_program_sheet(ws, fy_start):
    """
    Parse a per-branch Programs or Passive Programs sheet.
    Returns:
      bm_rows  : list of (month, activities, attendance) — Bookmobile summary
      data_rows: list of (month, branch, age_group, sessions, participation)
                 where one of sessions/participation may be None
    """
    bm_rows = []
    # data_rows keyed by (month, branch, age_group) → {'Sessions': v, 'Participation': v}
    data = {}

    # ── Bookmobile summary (rows 1-3) ────────────────────────────────────────
    bm_act = {}
    bm_att = {}
    if ws.nrows > 3:
        for col in range(MONTH_COL_START, MONTH_COL_START + 12):
            month = month_from_col(col)
            if not month:
                continue
            label_row1 = str(ws.cell_value(2, 1)).strip()
            label_row2 = str(ws.cell_value(3, 1)).strip()
            if 'Activities' in label_row1:
                bm_act[month] = cell(ws, 2, col)
            if 'Attendance' in label_row2:
                bm_att[month] = cell(ws, 3, col)

    for month in MONTHS:
        bm_rows.append((month, bm_act.get(month), bm_att.get(month)))

    # ── Main data sections ───────────────────────────────────────────────────
    current_age  = None
    current_stat = None

    for r in range(1, ws.nrows):
        col0 = str(ws.cell_value(r, 0)).strip()
        col1 = str(ws.cell_value(r, 1)).strip()

        if col0 and col0 not in ('', 'TOTAL'):
            parsed = parse_section_label(col0)
            if parsed:
                current_age, current_stat = parsed
            else:
                current_age = current_stat = None
            continue

        if current_age is None:
            continue
        if not col1 or col1 == 'TOTAL' or col1 == 'Activities' or col1 == 'Attendance':
            continue

        branch = BRANCH_RENAMES.get(col1, col1)

        for col in range(MONTH_COL_START, MONTH_COL_START + 12):
            month = month_from_col(col)
            if not month:
                continue
            v = cell(ws, r, col)
            key = (month, branch, current_age)
            if key not in data:
                data[key] = {'Sessions': None, 'Participation': None}
            data[key][current_stat] = v

    data_rows = []
    for month in MONTHS:
        for branch in ['Rock Hill', 'Clover', 'Fort Mill', 'Lake Wylie', 'York', 'Bookmobile/Outreach']:
            for age in ['0-5', '6-11', '12-18', 'Mixed Ages', '18+']:
                key = (month, branch, age)
                if key in data:
                    d = data[key]
                    data_rows.append((month, branch, age, d['Sessions'], d['Participation']))

    return bm_rows, data_rows


# ─────────────────────────────────────────────────────────────────────────────
# Parser: per-branch Training sheets (FY16-17 through FY18-19)
# ─────────────────────────────────────────────────────────────────────────────

def parse_branch_training_sheet(ws):
    """
    Returns list of (month, branch, staff_sessions, staff_participants, staff_hours).
    Only captures the first three sections (Staff Trained Sessions, Staff Participation,
    Staff Total Hours). Age-group public training sections are skipped (mostly zeros).
    """
    STAFF_SECTIONS = {
        'Staff Trained Sessions': 'staff_sessions',
        'Staff Participation':    'staff_participants',
        'Staff Total Hours':      'staff_hours',
    }

    data = {}
    current_stat = None

    for r in range(ws.nrows):
        col0 = str(ws.cell_value(r, 0)).strip()
        col1 = str(ws.cell_value(r, 1)).strip()

        # Stop at first age-group training section (public training)
        if col0.startswith('Age ') and 'Trained' in col0:
            break

        if col0 in STAFF_SECTIONS:
            current_stat = STAFF_SECTIONS[col0]
            continue

        if current_stat is None:
            continue
        if not col1 or col1 == 'TOTAL':
            continue

        branch = BRANCH_RENAMES.get(col1, col1)

        for col in range(MONTH_COL_START, MONTH_COL_START + 12):
            month = month_from_col(col)
            if not month:
                continue
            v = cell(ws, r, col)
            key = (month, branch)
            if key not in data:
                data[key] = {'staff_sessions': None, 'staff_participants': None, 'staff_hours': None}
            data[key][current_stat] = v

    rows = []
    for month in MONTHS:
        for branch in ['Rock Hill', 'Clover', 'Fort Mill', 'Lake Wylie', 'York', 'Bookmobile/Outreach']:
            key = (month, branch)
            if key in data:
                d = data[key]
                rows.append((month, branch, d['staff_sessions'], d['staff_participants'], d['staff_hours']))

    return rows


# ─────────────────────────────────────────────────────────────────────────────
# Parser: FY15-16 p.2  (system-wide, paired-column format)
# ─────────────────────────────────────────────────────────────────────────────

def parse_fy1516_programs(ws):
    """
    Returns:
      prog_rows:  list of (month, location_type, age_group, sessions, attendance)
      train_rows: list of (month, type, sessions, participants, hours)
    """
    # Months are in paired columns starting at col 3:
    #   col 3 = July #, col 4 = July Att., col 5 = Aug #, ... col 26 = June Att.
    PROG_MONTHS_COL = {m: 3 + i * 2 for i, m in enumerate(MONTHS)}

    LOCATION_SECTIONS = {
        'WITHIN THE LIBRARY': 'Within Library',
        'OUT OF THE LIBRARY':  'Out of Library',
        'CLASS VISITS':        'Class Visits',
    }
    AGE_GROUPS_1516 = {
        'Pre-Walkers': 'Pre-Walkers (0-12 mo)',
        'Walkers':     'Walkers (13-23 mo)',
        'Toddlers':    'Toddlers (24-35 mo)',
        'Preschool':   'Preschool (3-5 yr)',
        'Elem':        'Elem/Middle (6-11 yr)',
        'Teenage':     'Teenage (12-17 yr)',
        'Adult':       'Adult',
    }

    prog_rows  = []
    train_rows = []

    current_loc  = None
    in_training  = False

    # Training layout starts at the TRAINING STATISTICS row
    # Staff: rows 37-40 (Sessions, Number Trained, Hours of Training)
    # Public: rows 43-46 (Sessions, Number Trained, Hours of Training)
    # Training months: col 1 = July, col 2 = Aug, ... col 12 = June  (row 37 is header)
    TRAIN_MONTH_COL = {m: 1 + i for i, m in enumerate(MONTHS)}

    for r in range(ws.nrows):
        col0 = str(ws.cell_value(r, 0)).strip()

        # ── Training section ──────────────────────────────────────────────
        if 'TRAINING STATISTICS' in col0:
            in_training = True
            continue

        if in_training:
            # Check for STAFF / Public headers
            if col0 in ('STAFF', 'Public'):
                current_train_type = col0
                continue
            if col0 in ('Sessions', 'Number Trained', 'Hours of Training'):
                stat_map = {
                    'Sessions':           'sessions',
                    'Number Trained':     'participants',
                    'Hours of Training':  'hours',
                }
                stat = stat_map[col0]
                row_vals = {m: cell(ws, r, TRAIN_MONTH_COL[m]) for m in MONTHS}
                for m in MONTHS:
                    # Find or create entry
                    entry = next((x for x in train_rows if x[0] == m and x[1] == current_train_type), None)
                    if entry is None:
                        entry = [m, current_train_type, None, None, None]
                        train_rows.append(entry)
                    idx = {'sessions': 2, 'participants': 3, 'hours': 4}[stat]
                    entry[idx] = row_vals[m]
            continue

        # ── Program sections ──────────────────────────────────────────────
        if col0 in LOCATION_SECTIONS:
            current_loc = LOCATION_SECTIONS[col0]
            continue

        if current_loc is None:
            continue

        # Skip header and TOTAL rows
        if col0 in ('TOTAL', 'GRAND TOTALS', 'Technology Center ') or col0 == '':
            continue
        if col0 in ('', '#', 'Att.'):
            continue

        # Match age group
        age_label = None
        for prefix, label in AGE_GROUPS_1516.items():
            if col0.startswith(prefix):
                age_label = label
                break
        if age_label is None:
            continue

        for month in MONTHS:
            sc = PROG_MONTHS_COL[month]       # sessions col
            ac = sc + 1                        # attendance col
            if sc >= ws.ncols:
                continue
            sessions   = cell(ws, r, sc)
            attendance = cell(ws, r, ac) if ac < ws.ncols else None
            prog_rows.append((month, current_loc, age_label, sessions, attendance))

    return prog_rows, [tuple(x) for x in train_rows]


# ─────────────────────────────────────────────────────────────────────────────
# Spreadsheet styling helpers
# ─────────────────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill('solid', fgColor='1F4E79')
SUBHDR_FILL = PatternFill('solid', fgColor='2F75B6')
HEADER_FONT = Font(bold=True, color='FFFFFF')
BOLD        = Font(bold=True)

def write_header(ws, row, cols, fill=HEADER_FILL):
    for c, val in enumerate(cols, 1):
        cell_ = ws.cell(row=row, column=c, value=val)
        cell_.font = HEADER_FONT
        cell_.fill = fill
        cell_.alignment = Alignment(horizontal='center')


def auto_width(ws, min_w=8, max_w=30):
    for col in ws.columns:
        width = min_w
        for c in col:
            if c.value:
                width = max(width, min(len(str(c.value)) + 2, max_w))
        ws.column_dimensions[get_column_letter(col[0].column)].width = width


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

def main():
    import os
    os.makedirs("Data files/historical_data", exist_ok=True)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)   # remove default blank sheet

    # ── Sheet 1: Programs FY16-19 ────────────────────────────────────────────
    ws_prog = wb.create_sheet("Programs FY16-19")
    write_header(ws_prog, 1, ['FY', 'Month', 'Branch', 'Age Group', 'Sessions', 'Participation'])
    ws_prog.freeze_panes = 'A2'

    prog_sources = [
        ("Annual stats FY 15-16-17.xls", "FY16-17 Programs",  2016),
        ("Annual stats FY 17-18.xls",     "FY17-18 Programs",  2017),
        ("Annual stats FY 18-19.xls",     "FY18-19 Programs",  2018),
    ]
    bm_all = []   # accumulate for Bookmobile sheet
    row = 2
    for fname, sheetname, fy_start in prog_sources:
        wb_src = xlrd.open_workbook(BASE + fname)
        ws_src = wb_src.sheet_by_name(sheetname)
        bm_rows, data_rows = parse_branch_program_sheet(ws_src, fy_start)
        fy = fy_label(fy_start)
        for bm in bm_rows:
            bm_all.append((fy,) + bm)
        for r_data in data_rows:
            ws_prog.cell(row=row, column=1, value=fy)
            for c, v in enumerate(r_data, 2):
                ws_prog.cell(row=row, column=c, value=v)
            row += 1

    auto_width(ws_prog)

    # ── Sheet 2: Bookmobile Activities ──────────────────────────────────────
    ws_bm = wb.create_sheet("Bookmobile Activities")
    write_header(ws_bm, 1, ['FY', 'Month', 'Activities', 'Attendance'])
    ws_bm.freeze_panes = 'A2'
    for i, r_data in enumerate(bm_all, 2):
        for c, v in enumerate(r_data, 1):
            ws_bm.cell(row=i, column=c, value=v)
    auto_width(ws_bm)

    # ── Sheet 3: Passive Programs FY17-19 ────────────────────────────────────
    ws_pass = wb.create_sheet("Passive Programs FY17-19")
    write_header(ws_pass, 1, ['FY', 'Month', 'Branch', 'Age Group', 'Sessions', 'Participation'])
    ws_pass.freeze_panes = 'A2'

    passive_sources = [
        ("Annual stats FY 17-18.xls", "FY17-18 PASSIVE Programs", 2017),
        ("Annual stats FY 18-19.xls", "FY18-19 PASSIVE Programs", 2018),
    ]
    row = 2
    for fname, sheetname, fy_start in passive_sources:
        wb_src = xlrd.open_workbook(BASE + fname)
        ws_src = wb_src.sheet_by_name(sheetname)
        _, data_rows = parse_branch_program_sheet(ws_src, fy_start)
        fy = fy_label(fy_start)
        for r_data in data_rows:
            ws_pass.cell(row=row, column=1, value=fy)
            for c, v in enumerate(r_data, 2):
                ws_pass.cell(row=row, column=c, value=v)
            row += 1
    auto_width(ws_pass)

    # ── Sheet 4: Training FY16-19 ────────────────────────────────────────────
    ws_train = wb.create_sheet("Training FY16-19")
    write_header(ws_train, 1, ['FY', 'Month', 'Branch', 'Staff Sessions', 'Staff Participants', 'Staff Hours'])
    ws_train.freeze_panes = 'A2'

    train_sources = [
        ("Annual stats FY 15-16-17.xls", "FY16-17 Training", 2016),
        ("Annual stats FY 17-18.xls",     "FY17-18 Training", 2017),
        ("Annual stats FY 18-19.xls",     "FY18-19 Training", 2018),
    ]
    row = 2
    for fname, sheetname, fy_start in train_sources:
        wb_src = xlrd.open_workbook(BASE + fname)
        ws_src = wb_src.sheet_by_name(sheetname)
        data_rows = parse_branch_training_sheet(ws_src)
        fy = fy_label(fy_start)
        for r_data in data_rows:
            ws_train.cell(row=row, column=1, value=fy)
            for c, v in enumerate(r_data, 2):
                ws_train.cell(row=row, column=c, value=v)
            row += 1
    auto_width(ws_train)

    # ── Sheet 5: FY15-16 Programs (system-wide) ──────────────────────────────
    ws_1516 = wb.create_sheet("FY15-16 Programs")
    write_header(ws_1516, 1, ['Month', 'Location Type', 'Age Group', 'Sessions', 'Attendance'])
    ws_1516.freeze_panes = 'A2'

    wb_src = xlrd.open_workbook(BASE + "Annual stats FY 15-16-17.xls")
    ws_src = wb_src.sheet_by_name("FY15-16 p.2")
    prog_rows_1516, train_rows_1516 = parse_fy1516_programs(ws_src)
    for i, r_data in enumerate(prog_rows_1516, 2):
        for c, v in enumerate(r_data, 1):
            ws_1516.cell(row=i, column=c, value=v)
    auto_width(ws_1516)

    # ── Sheet 6: FY15-16 Training (system-wide) ──────────────────────────────
    ws_1516t = wb.create_sheet("FY15-16 Training")
    write_header(ws_1516t, 1, ['Month', 'Type', 'Sessions', 'Participants', 'Hours'])
    ws_1516t.freeze_panes = 'A2'
    for i, r_data in enumerate(train_rows_1516, 2):
        for c, v in enumerate(r_data, 1):
            ws_1516t.cell(row=i, column=c, value=v)
    auto_width(ws_1516t)

    wb.save(OUT)
    print(f"Saved: {OUT}")

    # Print summary
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        print(f"  {sheet}: {ws.max_row - 1} data rows")


if __name__ == '__main__':
    main()
