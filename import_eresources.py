"""
Import Monthly eResources usage from the director's tracking workbook
(e.g. 'Monthly Stats 2025-2026.xlsx') into the databases / usage_monthly tables.

Each vendor tab has its own layout (headers vary, some tabs have a second
table further down, Overdrive/Libby has two side-by-side tables), so sheets
are read by explicit (row, column) position rather than by column-name
lookup. The vendor -> database mapping was worked out with the library
director tab by tab and mirrors the source-sheet comments in
seed_eresource_databases() (seed_data.py) — that's the place to look if a
sheet's layout changes and this needs updating.

DISCUS is intentionally not handled here — that tab is blank in the
source workbook (seeded databases are inactive placeholders).

Usage (CLI):
    python import_eresources.py "path/to/Monthly Stats 2025-2026.xlsx" [fiscal_start_year]

Fiscal year runs July - June; fiscal_start_year is the calendar year the FY
begins in (e.g. 2025 for FY2026, "2025-2026"). If omitted, it's parsed from
a YYYY-YYYY pattern in the filename.
"""

import os
import re
import sys
from datetime import datetime

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from app import app, db
from models import EresourceDatabase, UsageMonthly

_MONTH_NAMES = {
    'january': 1, 'february': 2, 'march': 3, 'april': 4, 'may': 5, 'june': 6,
    'july': 7, 'august': 8, 'september': 9, 'october': 10, 'november': 11, 'december': 12,
    'feburary': 2,  # typo that appears in several vendor tabs
}

_STOP_LABELS = ('totals', 'total', 'average/month', 'average')


def parse_month(val):
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.month
    if isinstance(val, (int, float)):
        v = int(val)
        return v if 1 <= v <= 12 else None
    if isinstance(val, str):
        return _MONTH_NAMES.get(val.strip().lower())
    return None


def fiscal_year_for_month(month, fiscal_start_year):
    """Jul-Dec belongs to fiscal_start_year; Jan-Jun belongs to fiscal_start_year + 1."""
    return fiscal_start_year if month >= 7 else fiscal_start_year + 1


def _to_int(value):
    """Coerce a numeric-looking cell to int; None for blank/non-numeric cells."""
    if value is None or value == '':
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


# ── Sheet layouts ────────────────────────────────────────────────────────────
# Each block: (sheet_name, month_col, data_start_row, max_rows, {column: db_name})
# max_rows caps how far down we scan before giving up — the loop also stops
# early as soon as it hits a "Totals"/"Average" row in the month column, so
# this is just a safety ceiling, not an exact row count.

SHEET_BLOCKS = [
    ('ABC Mouse', 1, 2, 14, {
        2: 'ABC Mouse - Visits',
        3: 'ABC Mouse - Learning Activities',
    }),
    ('Ask a Librarian', 1, 2, 14, {
        2: 'Ask a Librarian - Questions',
    }),
    ('BiblioBoard', 1, 2, 14, {
        2: 'BiblioBoard - Sessions',
        3: 'BiblioBoard - Successful Requests',
    }),
    ('Brainfuse', 1, 2, 14, {
        2: 'Brainfuse - Sessions',
        3: 'Brainfuse - Skill Surfer Usage',
        5: 'Brainfuse - Tutoring Sessions',
    }),
    ('Data AxleReference USA', 1, 2, 14, {
        2: 'Data Axle/Reference USA - Sessions (Logins)',
        3: 'Data Axle/Reference USA - Searches',
        4: 'Data Axle/Reference USA - Downloads',
    }),
    ('DigitalLearn', 1, 2, 14, {
        2: 'DigitalLearn - Sessions',
        3: 'DigitalLearn - Completed Courses',
    }),
    ('EBSCO Flipster', 1, 2, 14, {
        2: 'EBSCO Flipster - Searches',
        4: 'EBSCO Flipster - Online Views',
        5: 'EBSCO Flipster - Downloads',
    }),
    ('Gale Ebooks', 1, 2, 14, {
        2: 'Gale eBooks - Sessions',
        3: 'Gale eBooks - Searches',
        4: 'Gale eBooks - Retrievals',
    }),
    ('Gale Presents Udemy', 1, 2, 14, {
        2: 'Gale Presents Udemy - Active Users',
        3: 'Gale Presents Udemy - Courses Enrolled',
        4: 'Gale Presents Udemy - Courses Started',
        5: 'Gale Presents Udemy - Courses Completed',
    }),
    ('Infobase The Mailbox', 1, 2, 14, {
        2: 'Infobase: The Mailbox - Searches',
        3: 'Infobase: The Mailbox - Sessions',
        4: 'Infobase: The Mailbox - Printable Downloads',
    }),
    ('Kanopy', 1, 3, 16, {
        2: 'Kanopy - Visits/Sessions',
        3: 'Kanopy - Plays',
    }),
    ('LibraryAware Newsletters', 1, 2, 14, {
        2: 'LibraryAware Newsletters - Unique Opens',
        3: 'LibraryAware Newsletters - Current Subscribers',
    }),
    ('Lote4Kids', 1, 2, 14, {
        2: 'Lote4Kids - Stories Watched',
        3: 'Lote4Kids - Activities',
        5: 'Lote4Kids - Logins',
    }),
    ('Mango', 1, 3, 16, {
        2: 'Mango - Sessions',
        3: 'Mango - Total Uses',
        5: 'Mango - Little Pim Uses',
    }),
    # Hoopla main table
    ('Midwest Tapehoopla', 1, 2, 14, {
        2: 'Hoopla - eBooks Instant',
        3: 'Hoopla - eBooks Flex',
        4: 'Hoopla - eAudio Instant',
        5: 'Hoopla - eAudio Flex',
        7: 'Hoopla - TV',
        8: 'Hoopla - Comics',
        9: 'Hoopla - Movies',
        10: 'Hoopla - Music',
    }),
    # Hoopla's BingePass breakdown — a second table further down the same sheet
    ('Midwest Tapehoopla', 1, 19, 14, {
        2: 'Hoopla - BingePass (Comics & eBooks)',
        3: 'Hoopla - BingePass (Audio)',
        4: 'Hoopla - BingePass (Courses & Videos)',
        5: 'Hoopla - BingePass (Magazines)',
    }),
    ('Newsbank', 1, 2, 14, {
        2: 'Newsbank - Logins',
        3: 'Newsbank - Searches',
        4: 'Newsbank - Documents Viewed',
    }),
    # Overdrive/Libby: Libby side starts at column A, Sora side at column H
    ('OverdriveLibby', 1, 3, 16, {
        2: 'Overdrive/Libby - eBooks',
        3: 'Overdrive/Libby - eAudio',
        4: 'Overdrive/Libby - Magazines',
        5: 'Overdrive/Libby - Streaming',
    }),
    ('OverdriveLibby', 8, 3, 16, {
        11: 'Sora - Total',
    }),
    ('Proquest Ancestry Library Editi', 1, 2, 14, {
        2: 'Proquest: Ancestry Library Edition - Sessions',
        3: 'Proquest: Ancestry Library Edition - Searches',
    }),
    ('Proquest Fold3', 1, 2, 14, {
        2: 'Proquest: Fold3 - Searches',
        3: 'Proquest: Fold3 - Item Requests',
    }),
    ('Proquest HeritageQuest', 1, 2, 14, {
        2: 'Proquest: HeritageQuest - Sessions',
    }),
    ('Salem Press', 1, 2, 14, {
        2: 'Salem Press - Regular Search',
        3: 'Salem Press - Result Click',
    }),
    ('Tutor.com', 1, 2, 14, {
        2: 'Tutor.com - Sessions',
        3: 'Tutor.com - Past Session Views',
    }),
    ('Value Line', 1, 2, 14, {
        2: 'Value Line - Logins',
        4: 'Value Line - Downloads',
    }),
    ('Weiss Financial Services', 1, 2, 14, {
        3: 'Weiss Financial Services - Sessions',
        4: 'Weiss Financial Services - Searches',
        5: 'Weiss Financial Services - Users',
    }),
]


def import_block(ws, month_col, data_start_row, max_rows, col_to_db, db_lookup,
                  fiscal_start_year, warnings):
    created = updated = 0
    for r in range(data_start_row, data_start_row + max_rows):
        month_cell = ws.cell(row=r, column=month_col).value
        if isinstance(month_cell, str) and month_cell.strip().lower() in _STOP_LABELS:
            break
        month = parse_month(month_cell)
        if month is None:
            continue
        year = fiscal_year_for_month(month, fiscal_start_year)

        for col, db_name in col_to_db.items():
            database = db_lookup.get(db_name)
            if database is None:
                warnings.append(f"Unknown database '{db_name}' — check seed data")
                continue
            value = _to_int(ws.cell(row=r, column=col).value)
            if value is None:
                continue

            usage = UsageMonthly.query.filter_by(
                database_id=database.id, year=year, month=month
            ).first()
            if usage is None:
                usage = UsageMonthly(database_id=database.id, year=year, month=month)
                db.session.add(usage)
                created += 1
            else:
                updated += 1
            usage.usage_count = value

    return created, updated


def guess_fiscal_start_year(path):
    m = re.search(r'(\d{4})-(\d{4})', os.path.basename(path))
    return int(m.group(1)) if m else None


def do_import(wb, fiscal_start_year):
    """
    Import all recognised vendor sheets from an open openpyxl workbook
    (loaded with data_only=True). Must be called within an active Flask
    app context. Returns (created, updated, warnings).
    """
    db_lookup = {d.name: d for d in EresourceDatabase.query.all()}
    total_created = total_updated = 0
    warnings = []

    for sheet_name, month_col, data_start_row, max_rows, col_to_db in SHEET_BLOCKS:
        if sheet_name not in wb.sheetnames:
            warnings.append(f"Sheet '{sheet_name}' not found in workbook — skipped")
            continue
        c, u = import_block(wb[sheet_name], month_col, data_start_row, max_rows,
                             col_to_db, db_lookup, fiscal_start_year, warnings)
        total_created += c
        total_updated += u

    db.session.commit()
    return total_created, total_updated, warnings


def run(excel_path, fiscal_start_year=None):
    if not os.path.exists(excel_path):
        print(f'ERROR: Cannot find {excel_path}')
        sys.exit(1)

    fiscal_start_year = fiscal_start_year or guess_fiscal_start_year(excel_path)
    if not fiscal_start_year:
        print('ERROR: Could not infer fiscal start year from filename; pass it explicitly, '
              'e.g. python import_eresources.py "Monthly Stats 2025-2026.xlsx" 2025')
        sys.exit(1)

    print(f'Opening {excel_path} '
          f'(FY{fiscal_start_year + 1}, Jul {fiscal_start_year} - Jun {fiscal_start_year + 1}) ...')
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    with app.app_context():
        created, updated, warnings = do_import(wb, fiscal_start_year)

    print(f'\n{created} usage_monthly rows created, {updated} updated')
    for w in warnings:
        print(f'  Warning: {w}')
    print('\nDone!')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python import_eresources.py "path/to/Monthly Stats 2025-2026.xlsx" [fiscal_start_year]')
        sys.exit(1)
    path = sys.argv[1]
    fy = int(sys.argv[2]) if len(sys.argv) > 2 else None
    run(path, fy)
