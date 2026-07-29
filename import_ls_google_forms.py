"""
Import program/branch stats from all three Google Forms export files in Data files/Ls/.

Files processed (all share the 'Form Responses 1' sheet format):
  other_missing_stats.xlsx       – FY2024-25 bulk (222 rows)
  possibly_more_missing_stats.xlsx – FY2024-25 early months (29 rows)
  2025_missing_stats.xlsx        – FY2025-26 (114 rows)

Metrics written include ONSITE program sessions+attendance by age group,
OFFSITE/VIRTUAL totals, outreach, take-and-makes, door count, PC/WiFi,
ILLs, room use, 1-on-1, and staff training.

Registration metrics (SIRSI-sourced) are never overwritten.
Rows for 'YCL (System Wide)' are skipped.
For duplicate form submissions (same branch+month), the latest row wins.

Usage:
    python import_ls_google_forms.py           # dry run — shows what would change
    python import_ls_google_forms.py --apply   # write to DB
"""

import os
import sys
from dotenv import load_dotenv
load_dotenv()

import openpyxl
from flask import Flask
from models import db
from import_excel import (
    build_branch_lookup, build_metric_lookup,
    import_google_forms_stats,
)

app = Flask(__name__)
_db_url = os.environ.get('DATABASE_URL', 'sqlite:///librarystats.db')
if _db_url.startswith('postgres://'):
    _db_url = _db_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = _db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db.init_app(app)

FILES = [
    'Data files/Ls/other_missing_stats.xlsx',
    'Data files/Ls/possibly_more_missing_stats.xlsx',
    'Data files/Ls/2025_missing_stats.xlsx',
]


def main(dry_run=True):
    print(f"{'DRY RUN — ' if dry_run else ''}Importing Google Forms branch stats from Ls/\n")

    with app.app_context():
        branch_lookup = build_branch_lookup()
        bs_metrics, bs_cat = build_metric_lookup('Branch Stats')
        if not bs_cat:
            print("ERROR: Branch Stats category not found in DB")
            sys.exit(1)

        total_created = total_updated = 0

        for path in FILES:
            if not os.path.exists(path):
                print(f"  SKIP {path}: file not found")
                continue

            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            if 'Form Responses 1' not in wb.sheetnames:
                print(f"  SKIP {path}: no 'Form Responses 1' sheet")
                wb.close()
                continue

            ws = wb['Form Responses 1']
            print(f"  {os.path.basename(path)}")

            if dry_run:
                # Peek at what would be written without committing
                from import_excel import (
                    parse_month, col_index, GOOGLE_FORMS_STATS_MAP,
                )
                from datetime import datetime as dt
                rows = list(ws.iter_rows(values_only=True))
                headers = rows[0]
                ts_idx = col_index(headers, 'Timestamp')
                mo_idx = col_index(headers, 'Select Month')
                br_idx = col_index(headers, 'Select Branch')
                buckets = {}
                bad = 0
                for row in rows[1:]:
                    if all(v is None for v in row):
                        continue
                    ts = row[ts_idx]
                    mo = parse_month(row[mo_idx])
                    br = row[br_idx]
                    if not isinstance(ts, dt) or not mo or not br:
                        bad += 1
                        continue
                    br = str(br).strip()
                    if 'system wide' in br.lower():
                        continue
                    branch = branch_lookup.get(br) or branch_lookup.get(br.lower())
                    if branch is None:
                        continue
                    yr = ts.year - 1 if mo > ts.month else ts.year
                    buckets[(yr, mo, branch.name)] = True
                print(f"    Would write to {len(buckets)} branch/month entries "
                      f"({bad} rows skipped)")
                for (yr, mo, bname) in sorted(buckets):
                    print(f"      {bname:25s} {yr}-{mo:02d}")
            else:
                created, updated, periods, warnings = import_google_forms_stats(
                    ws, bs_cat, bs_metrics, branch_lookup
                )
                print(f"    created={created}  updated={updated}  periods={len(periods)}")
                for w in warnings:
                    print(f"    WARNING: {w}")
                total_created += created
                total_updated += updated

            wb.close()
            print()

        if not dry_run:
            print(f"Done. Total: {total_created} entries created, {total_updated} updated.")
        else:
            print("Dry run complete. Run with --apply to write to DB.")


if __name__ == '__main__':
    main(dry_run='--apply' not in sys.argv)
