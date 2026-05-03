"""
Patch FY24-25 programming gaps from '24-25 Annual Stats.xlsx' Program Comparison sheet.

Targets only the branch/month combinations that currently have zero programming data:
  - Clover       Dec 2024
  - Lake Wylie   Jul 2024
  - Bookmobile   Jul 2024, Aug 2024, Sep 2024, Jun 2025

Metric mapping:
  Ls row label                      → DB metric IDs
  Synchronus Age 0-5 Sessions       → 17 (ONSITE Sessions 0-5)
  Synchronus Age 6-11 Sessions      → 18 (ONSITE Sessions 6-11)
  Synchronus Age 12-18 Sessions     → 19 (ONSITE Sessions 12-18)
  Synchronus Age 19+ Sessions       → 20 (ONSITE Sessions 19+)
  Synchronus General Interest S.    → 21 (ONSITE Sessions General Interest)
  S. In-Person Offsite Sessions     → 28 (OFFSITE Sessions 6-11, matching existing convention)
  Age 0-5 Attendance                → 22 (ONSITE Attendance 0-5)
  Age 6-11 Attendance               → 23 (ONSITE Attendance 6-11)
  Age 12-18 Attendance              → 24 (ONSITE Attendance 12-18)
  Age 19+ Attendance                → 25 (ONSITE Attendance 19+)
  General Interest Attendance       → 26 (ONSITE Attendance General Interest)
  S. In-Person Offsite Attendance   → 33 (OFFSITE Attendance 6-11, matching existing convention)
"""

import os
import sys
from dotenv import load_dotenv
load_dotenv()

import openpyxl
from flask import Flask
from models import db, Entry, EntryValue, Branch
from sqlalchemy import and_

# ── App setup ─────────────────────────────────────────────────────────────────

app = Flask(__name__)
_db_url = os.environ.get('DATABASE_URL', 'sqlite:///librarystats.db')
if _db_url.startswith('postgres://'):
    _db_url = _db_url.replace('postgres://', 'postgresql://', 1)
app.config['SQLALCHEMY_DATABASE_URI'] = _db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db.init_app(app)

# ── Spreadsheet parsing ────────────────────────────────────────────────────────

XLSX = 'Data files/Ls/24-25 Annual Stats.xlsx'

# Column index → calendar month number (FY Jul-Jun)
COL_TO_MONTH = {3: 7, 4: 8, 5: 9, 6: 10, 7: 11, 8: 12,
                9: 1, 10: 2, 11: 3, 12: 4, 13: 5, 14: 6}
COL_TO_YEAR = {3: 2024, 4: 2024, 5: 2024, 6: 2024, 7: 2024, 8: 2024,
               9: 2025, 10: 2025, 11: 2025, 12: 2025, 13: 2025, 14: 2025}

BRANCH_ORDER = ['Rock Hill', 'Clover', 'Fort Mill', 'Lake Wylie', 'York', 'Bookmobile']

# Row ranges for each metric section (1-indexed, matching openpyxl row numbers)
# Each section has 6 data rows in branch order
SECTIONS = [
    # (first_data_row, metric_id)
    (16, 17),   # ONSITE Sessions 0-5
    (23, 18),   # ONSITE Sessions 6-11
    (30, 19),   # ONSITE Sessions 12-18
    (37, 20),   # ONSITE Sessions 19+
    (44, 21),   # ONSITE Sessions General Interest
    (58, 28),   # Offsite Sessions total → stored as OFFSITE 6-11 to match existing convention
    (77, 22),   # ONSITE Attendance 0-5
    (84, 23),   # ONSITE Attendance 6-11
    (91, 24),   # ONSITE Attendance 12-18
    (98, 25),   # ONSITE Attendance 19+
    (105, 26),  # ONSITE Attendance General Interest
    (119, 33),  # Offsite Attendance total → stored as OFFSITE Att. 6-11
]

def parse_sheet():
    """Returns dict: {(branch_name, year, month): {metric_id: value}}"""
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Program Comparison p.5']
    rows = {r: list(ws.iter_rows(min_row=r, max_row=r, values_only=True))[0]
            for r in range(1, ws.max_row + 1)}

    data = {}
    for first_row, metric_id in SECTIONS:
        for b_offset, branch_name in enumerate(BRANCH_ORDER):
            row = rows[first_row + b_offset]
            for col_idx, month in COL_TO_MONTH.items():
                year = COL_TO_YEAR[col_idx]
                raw = row[col_idx]
                val = int(raw) if isinstance(raw, (int, float)) and raw is not None else 0
                key = (branch_name, year, month)
                if key not in data:
                    data[key] = {}
                data[key][metric_id] = val
    return data

# ── DB helpers ─────────────────────────────────────────────────────────────────

BRANCH_NAME_MAP = {
    'Rock Hill':  'Rock Hill',
    'Clover':     'Clover',
    'Fort Mill':  'Fort Mill',
    'Lake Wylie': 'Lake Wylie',
    'York':       'York',
    'Bookmobile': 'Bookmobile/Outreach',
}

CATEGORY_ID = 1  # Branch Stats

PROG_METRIC_IDS = set(range(17, 47))

def has_programming(entry_id):
    """Return True if this entry already has any non-zero programming metric."""
    vals = EntryValue.query.filter(
        EntryValue.entry_id == entry_id,
        EntryValue.metric_id.in_(PROG_METRIC_IDS),
        EntryValue.value_number != None,
        EntryValue.value_number != 0
    ).first()
    return vals is not None

def upsert_value(entry_id, metric_id, value):
    ev = EntryValue.query.filter_by(entry_id=entry_id, metric_id=metric_id).first()
    if ev:
        ev.value_number = float(value)
    else:
        ev = EntryValue(entry_id=entry_id, metric_id=metric_id, value_number=float(value))
        db.session.add(ev)

# ── Main ───────────────────────────────────────────────────────────────────────

GAPS = [
    ('Clover',     2024, 12),
    ('Lake Wylie', 2024,  7),
    ('Bookmobile', 2024,  7),
    ('Bookmobile', 2024,  8),
    ('Bookmobile', 2024,  9),
    ('Bookmobile', 2025,  6),
]

def main(dry_run=False):
    print(f"{'DRY RUN — ' if dry_run else ''}Patching FY24-25 programming gaps\n")

    sheet_data = parse_sheet()

    with app.app_context():
        # Build branch_name → branch_id map
        branches = {b.name: b.id for b in Branch.query.all()}

        for (ls_branch, year, month) in GAPS:
            db_branch_name = BRANCH_NAME_MAP[ls_branch]
            branch_id = branches.get(db_branch_name)
            if not branch_id:
                print(f"  ERROR: branch '{db_branch_name}' not found in DB")
                continue

            entry = Entry.query.filter_by(
                category_id=CATEGORY_ID,
                branch_id=branch_id,
                year=year,
                month=month
            ).first()
            if not entry:
                print(f"  SKIP {ls_branch} {year}-{month:02d}: no entry found in DB")
                continue

            if has_programming(entry.id):
                print(f"  SKIP {ls_branch} {year}-{month:02d}: already has programming data (entry {entry.id})")
                continue

            key = (ls_branch, year, month)
            metrics = sheet_data.get(key, {})
            total_sessions = sum(metrics.get(m, 0) for m in [17,18,19,20,21,28])

            print(f"  PATCH {ls_branch} {year}-{month:02d} (entry {entry.id}): "
                  f"{total_sessions} sessions, "
                  f"{sum(metrics.get(m,0) for m in [22,23,24,25,26,33])} attendance")

            if not dry_run:
                for metric_id, value in metrics.items():
                    if value:  # skip zeros to keep the DB clean
                        upsert_value(entry.id, metric_id, value)
                db.session.commit()
                print(f"    → committed {sum(1 for v in metrics.values() if v)} non-zero values")

    if dry_run:
        print("\nDry run complete. Run with --apply to write to DB.")
    else:
        print("\nDone.")

if __name__ == '__main__':
    dry_run = '--apply' not in sys.argv
    main(dry_run=dry_run)
